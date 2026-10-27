from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path("docs/Credit_Risk_CICD_Report.docx")

def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); props.append(shd)

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers)); table.style = 'Table Grid'; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]; c.text = h; shade(c, '1F4E78'); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.color.rgb = __import__('docx').shared.RGBColor(255,255,255); r.bold = True
    for row in rows:
        cells=table.add_row().cells
        for i,value in enumerate(row): cells[i].text=str(value); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()

def main():
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
    styles=doc.styles; styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10.5)
    title=doc.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.add_run('Credit Risk CI CD Quality Gate')
    sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; sub.add_run('Project Report').italic=True
    def h(t): doc.add_heading(t, level=1)
    def p(t): doc.add_paragraph(t)
    h('Purpose and design')
    p('We built a reproducible credit-card default-risk workflow that blocks changes which fail tests or materially reduce model quality. Git versions source code and configuration, DVC versions the raw v1 and v2 datasets in DagsHub Storage, YAML controls behavior, MLflow records experiments, and GitHub Actions runs the quality gate on pull requests.')
    h('Data and feature engineering')
    p('The target is High_Risk_Customer, with approximately 19 percent positive cases. Customer_ID is excluded from predictors. Numeric values are median-imputed and categorical values are most-frequent-imputed within the fitted training pipeline, preventing test-set leakage. The reusable feature function creates utilization, a high-utilization flag, payment-delay band, and monthly-spend-to-limit. Unit tests cover expected values and a zero-credit-limit edge case.')
    add_table(doc,['Feature','Definition','Reason'],[['Credit Utilization','Current Balance / Credit Limit','Repayment stress indicator'],['High Utilization Flag','Utilization at or above YAML cutoff','Non-linear risk boundary'],['Payment Delay Band','On time, short, or long delay','Separates persistent delinquency'],['Monthly Spend to Limit','Monthly Spend / Credit Limit','Capacity pressure indicator']])
    h('Training and model selection')
    p('We compared Logistic Regression and Random Forest across five random seeds, producing ten tracked runs. We selected PR-AUC as the primary metric because accuracy is misleading for an imbalanced target: a mostly low-risk prediction can have high accuracy while missing risky customers. Logistic Regression produced mean PR-AUC 0.8129 with sample standard deviation 0.0220, compared with Random Forest at 0.8033 and 0.0253. The registered candidate is Logistic Regression with class_weight balanced, C 1.0, and max_iter 1000.')
    h('Quality gate and data change drill')
    p('The CI gate runs pytest, trains with configs/ci.yaml, and compares the candidate to baselines/approved.json. It requires PR-AUC no more than 0.01 below baseline and recall at 50 percent precision of at least 0.45. The baseline may be changed only through a reviewed promotion pull request with linked experiment evidence.')
    add_table(doc,['Version','Mean PR AUC','Sample SD','Decision'],[['v1','0.8112','0.0256','Baseline'],['v2','0.8483','0.0060','Accept']])
    p('Across seeds 42, 43, and 44, v2 improved mean PR-AUC by 0.0371. All three v2 runs passed each candidate rule: strict no-drop, recommended minus 0.01, and permissive minus 0.03. We therefore accept v2, while continuing the same gate and monitoring process.')
    h('Fairness and conclusion')
    p('For the final v1 random-forest fairness check at threshold 0.5, false-negative rates were 20.0 percent for High income, 17.6 percent for Low income, and 12.5 percent for Medium income. The 7.5 percentage-point range warrants monitoring, but supports are modest, so it is not conclusive. The model is conditionally acceptable for this classroom project; we would revisit the decision if the gap persists or widens with more data.')
    add_table(doc,['Income Bracket','Support','False Negative Rate','False Positive Rate'],[['High','42','0.200','0.135'],['Low','83','0.176','0.121'],['Medium','75','0.125','0.136']])
    doc.save(OUT)
if __name__=='__main__': main()
